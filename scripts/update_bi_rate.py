
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from io import BytesIO

URL = "https://www.bi.go.id/id/statistik/indikator/bi-rate.aspx"
ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "data" / "bi_rate.json"


def main():
    session = requests.Session()
    session.headers.update({
        "User-Agent": "Mozilla/5.0 (compatible; BI-Rate-DataBot/1.0)"
    })

    # get halaman untuk memperoleh state form terbaru.
    response = session.get(URL, timeout=30)
    response.raise_for_status()

    soup = BeautifulSoup(response.text, "html.parser")
    form = soup.find("form")
    if not form:
        raise RuntimeError("Form BI Rate tidak ditemukan.")

    payload = {}
    for field in form.select("input[name]"):
        name = field.get("name")
        if name:
            payload[name] = field.get("value", "")

    payload["__EVENTTARGET"] = (
        "ctl00$ctl51$g_3800da0f_298c_4f6f_a7c2_201434584198$ctl00$ButtonExport"
    )
    payload["__EVENTARGUMENT"] = ""

    # Kirim POST untuk meminta file ekspor.
    export = session.post(URL, data=payload, timeout=60)
    export.raise_for_status()

    content_type = export.headers.get("Content-Type", "").lower()
    if "spreadsheetml" not in content_type and not export.content.startswith(b"PK"):
        print("Status:", export.status_code)
        print("Content-Type:", content_type)
        print("URL akhir:", export.url)
        print("Awal respons:", export.text[:1000])
    raise RuntimeError("Respons bukan XLSX; lihat log di atas.")
        )

    # Baca workbook Excel.
    workbook = load_workbook(BytesIO(export.content), data_only=True, read_only=True)
    sheet = workbook.active

    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    # Tampilkan beberapa baris untuk memvalidasi struktur workbook.
    print("Jumlah baris Excel:", len(rows))
    for row in rows[:8]:
        print(repr(row))

    # Belum menulis JSON sampai struktur kolom dipastikan.
    print("Unduhan XLSX berhasil. Struktur kolom perlu diverifikasi sebelum sinkronisasi.")


if __name__ == "__main__":
    main()

